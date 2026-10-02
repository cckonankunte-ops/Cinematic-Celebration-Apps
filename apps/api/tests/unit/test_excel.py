"""Unit tests for the ``.xlsx`` export builders (pure, no DB).

Build a workbook from a couple of rows, load it back from the ``BytesIO`` with
openpyxl, and assert the header row plus a Day Summary footer where expected.
"""

from __future__ import annotations

from datetime import datetime, timezone

from openpyxl import load_workbook

from app.schemas.analytics import (
    AdvanceAmountRow,
    AdvanceAmountsResult,
    AnalyticsDaySummary,
    SlotAmountRow,
    SlotAmountsResult,
)
from app.schemas.sheets import CakesheetRow, TimesheetRow
from app.services import excel as excel_service

_NOW = datetime(2025, 1, 1, 10, 30, tzinfo=timezone.utc)


def _timesheet_rows() -> list[TimesheetRow]:
    return [
        TimesheetRow(
            slot_id=1,
            slot_sort_order=1,
            slot_description="9 AM - 12 PM",
            plan_title="Starlight",
            booking_name="Priya",
            phone="9876543210",
            people=6,
            name_on_cake="Happy Birthday",
            cake="Chocolate truffle",
            special_decor="Rose bouquet",
            combos="Party combo",
            occasion="Birthday - Ravi",
            amount_due_paise=400000,
        )
    ]


def test_timesheet_xlsx_has_header_and_rows() -> None:
    buffer = excel_service.build_timesheet_xlsx(_timesheet_rows())
    wb = load_workbook(buffer)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header[0] == "Slot"
    assert "Amount Due (Rs.)" in header
    # One header row + one data row.
    assert ws.max_row == 2
    # Money column is shown in rupees (paise / 100).
    assert ws.cell(row=2, column=header.index("Amount Due (Rs.)") + 1).value == 4000.0


def test_cakesheet_xlsx_structure() -> None:
    rows = [
        CakesheetRow(
            slot_id=1,
            slot_sort_order=1,
            slot_description="9 AM - 12 PM",
            cake="Chocolate truffle",
            name_on_cake="Happy Birthday",
        )
    ]
    buffer = excel_service.build_cakesheet_xlsx(rows)
    wb = load_workbook(buffer)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == ["Slot", "Cake", "Name on Cake"]
    assert ws.max_row == 2


def test_slot_amounts_xlsx_has_day_summary_footer() -> None:
    result = SlotAmountsResult(
        rows=[
            SlotAmountRow(
                booking_name="Priya",
                booking_reference="CC-ABC",
                total_paise=400000,
                amount_paid_paise=50000,
                today_income_paise=50000,
                amount_due_paise=350000,
                discount_paise=0,
                food_paise=0,
                cleaning_paise=0,
                other_paise=0,
                created_at=_NOW,
            )
        ],
        summary=AnalyticsDaySummary(
            total_paise=400000,
            amount_paid_paise=50000,
            amount_due_paise=350000,
            discount_paise=0,
            food_paise=0,
            cleaning_paise=0,
            other_paise=0,
        ),
    )
    buffer = excel_service.build_slot_amounts_xlsx(result)
    wb = load_workbook(buffer)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header[0] == "Booking Name"
    # Header + 1 data row + Day Summary footer.
    assert ws.max_row == 3
    assert ws.cell(row=3, column=1).value == "Day Summary"
    assert ws.cell(row=3, column=3).value == 4000.0


def test_advance_amounts_xlsx_has_day_summary_footer() -> None:
    result = AdvanceAmountsResult(
        rows=[
            AdvanceAmountRow(
                booking_name="Priya",
                booking_reference="CC-ABC",
                total_paise=400000,
                amount_paid_paise=50000,
                created_at=_NOW,
            )
        ],
        summary=AnalyticsDaySummary(
            total_paise=400000,
            amount_paid_paise=50000,
            amount_due_paise=350000,
            discount_paise=0,
            food_paise=0,
            cleaning_paise=0,
            other_paise=0,
        ),
    )
    buffer = excel_service.build_advance_amounts_xlsx(result)
    wb = load_workbook(buffer)
    ws = wb.active
    header = [cell.value for cell in ws[1]]
    assert header == ["Booking Name", "Reference", "Total (Rs.)", "Paid (Rs.)"]
    assert ws.max_row == 3
    assert ws.cell(row=3, column=1).value == "Day Summary"
